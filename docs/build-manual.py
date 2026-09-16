"""
Generates AuthMicroservice-Manual.docx next to this script.

Run:
    py docs/build-manual.py

Requires:
    py -m pip install --user python-docx
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUTPUT = Path(__file__).with_name("AuthMicroservice-Manual.docx")

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
CODE_BG = "F2F2F2"
BODY_FONT = "Calibri"
CODE_FONT = "Consolas"


def _shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _cell_borders(cell, size="4", color="BFBFBF"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)
        borders.append(el)
    tc_pr.append(borders)


def add_code_block(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Cm(16)
    cell = table.rows[0].cells[0]
    cell.width = Cm(16)
    _shade(cell, CODE_BG)
    _cell_borders(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    cell.text = ""
    for i, line in enumerate(code.rstrip("\n").split("\n")):
        para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run(line if line else " ")
        run.font.name = CODE_FONT
        run.font.size = Pt(9.5)
    doc.add_paragraph()  # trailing spacer


def add_inline_code(paragraph, text):
    run = paragraph.add_run(text)
    run.font.name = CODE_FONT
    run.font.size = Pt(10)
    return run


def add_bullet(doc, text, code_terms=()):
    para = doc.add_paragraph(style="List Bullet")
    _add_mixed_runs(para, text, code_terms)


def _add_mixed_runs(paragraph, text, code_terms):
    if not code_terms:
        paragraph.add_run(text)
        return
    remaining = text
    while remaining:
        first_idx = -1
        first_term = None
        for term in code_terms:
            idx = remaining.find(term)
            if idx != -1 and (first_idx == -1 or idx < first_idx):
                first_idx = idx
                first_term = term
        if first_idx == -1:
            paragraph.add_run(remaining)
            break
        if first_idx > 0:
            paragraph.add_run(remaining[:first_idx])
        add_inline_code(paragraph, first_term)
        remaining = remaining[first_idx + len(first_term):]


def add_para(doc, text, code_terms=()):
    para = doc.add_paragraph()
    _add_mixed_runs(para, text, code_terms)
    return para


def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = ACCENT
        run.font.name = BODY_FONT
    return heading


def add_table(doc, headers, rows, col_widths_cm=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = ""
        para = hdr[i].paragraphs[0]
        run = para.add_run(text)
        run.bold = True
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.rows[r].cells[c]
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.add_run(str(val))
    if col_widths_cm:
        for col, w in zip(table.columns, col_widths_cm):
            for cell in col.cells:
                cell.width = Cm(w)
    doc.add_paragraph()
    return table


def insert_toc(doc):
    para = doc.add_paragraph()
    run = para.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    fld_char3 = OxmlElement("w:t")
    fld_char3.text = "Right-click and choose 'Update Field' in Word to populate."
    fld_char4 = OxmlElement("w:fldChar")
    fld_char4.set(qn("w:fldCharType"), "end")
    r_elem = run._r
    r_elem.append(fld_char1)
    r_elem.append(instr)
    r_elem.append(fld_char2)
    r_elem.append(fld_char3)
    r_elem.append(fld_char4)


def build_document():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)

    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    # ---------------- Cover ----------------
    cover_title = doc.add_paragraph()
    cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_title.paragraph_format.space_before = Pt(120)
    tr = cover_title.add_run("AuthMicroservice")
    tr.bold = True
    tr.font.size = Pt(36)
    tr.font.color.rgb = ACCENT
    tr.font.name = BODY_FONT

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("ASP.NET Core 8 Authentication Library and Microservice")
    sr.italic = True
    sr.font.size = Pt(16)
    sr.font.name = BODY_FONT

    umanual = doc.add_paragraph()
    umanual.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ur = umanual.add_run("User Manual")
    ur.font.size = Pt(20)
    ur.font.name = BODY_FONT
    ur.font.color.rgb = ACCENT

    for _ in range(2):
        doc.add_paragraph()

    meta_lines = [
        ("Version", "1.0"),
        ("Release date", "2026-09-10"),
        ("Author", "Synergy Software"),
        ("License", "MIT"),
    ]
    for label, value in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lr = p.add_run(f"{label}: ")
        lr.bold = True
        p.add_run(value)

    doc.add_page_break()

    # ---------------- TOC ----------------
    add_heading(doc, "Table of Contents", level=1)
    add_para(
        doc,
        "When you open this document in Microsoft Word, right-click the field below "
        "and select 'Update Field' (or press F9) to populate the table of contents. "
        "Word Online and LibreOffice will also render it automatically.",
    )
    insert_toc(doc)
    doc.add_page_break()

    # ---------------- 1. Introduction ----------------
    add_heading(doc, "1. Introduction", level=1)
    add_para(
        doc,
        "AuthMicroservice is a reusable authentication component for ASP.NET Core 8. "
        "It can be consumed in two ways: as a plug-in library referenced from your own "
        "Web API project, or as a standalone containerised microservice deployed next to "
        "your application. It is similar in spirit to Keycloak but intentionally lightweight, "
        "idiomatic to .NET, and configured entirely through appsettings.json and environment "
        "variables.",
    )
    add_para(
        doc,
        "The goal of the project is to reduce the time-to-first-authentication for new "
        ".NET services from days to minutes. A single call to AddAuthMicroservice() in "
        "Program.cs is enough to register users, log them in, issue and rotate JWT access "
        "tokens, verify email addresses, reset passwords, enforce lockout, and expose the "
        "authenticated identity as roles and claims.",
        code_terms=("AddAuthMicroservice()", "Program.cs", "appsettings.json"),
    )

    # ---------------- 2. Feature overview ----------------
    add_heading(doc, "2. Feature Overview", level=1)
    features = [
        ("Registration, login, and logout using email and password.", ()),
        ("JWT access tokens with refresh-token rotation. Refresh tokens are stored as "
         "SHA-256 hashes and revoked on replay.", ()),
        ("Email verification links and password-reset links delivered over SMTP through "
         "MailKit. The transport is pluggable via IEmailSender.", ("IEmailSender",)),
        ("Change-password endpoint that revokes all existing refresh tokens for safety.", ()),
        ("Account lockout after a configurable number of failed attempts.", ()),
        ("Roles and custom claims backed by ASP.NET Core Identity.", ()),
        ("Provider-agnostic Entity Framework Core: pick SqlServer, Postgres, Sqlite, or "
         "InMemory in appsettings.json.", ("appsettings.json",)),
        ("All configuration comes from appsettings.json and environment variables. No "
         "secrets are hard-coded, and startup fails fast on missing or invalid config.",
         ("appsettings.json",)),
        ("Swagger UI, a production Dockerfile, docker-compose with Mailhog for local "
         "email preview, a sample consumer project, plus unit and integration test "
         "suites.", ("docker-compose",)),
    ]
    for text, terms in features:
        add_bullet(doc, text, terms)

    # ---------------- 3. Solution architecture ----------------
    add_heading(doc, "3. Solution Architecture", level=1)
    add_para(
        doc,
        "The solution is organised so that everything reusable lives in a single core "
        "library, while database-provider-specific code lives in dedicated migration "
        "assemblies. A standalone host and a sample consumer wire the library together "
        "in two typical deployment styles.",
    )
    add_table(
        doc,
        headers=["Project", "Purpose"],
        rows=[
            ["AuthMicroservice.Core",
             "Reusable library: domain, EF context, JWT + refresh + email services, "
             "minimal-API endpoints, DI extensions."],
            ["AuthMicroservice.Migrations.SqlServer",
             "SQL Server-specific EF Core migrations."],
            ["AuthMicroservice.Migrations.Postgres",
             "PostgreSQL-specific EF Core migrations."],
            ["AuthMicroservice.Migrations.Sqlite",
             "SQLite-specific EF Core migrations."],
            ["AuthMicroservice.Api",
             "Standalone Web API host. This is the image built by the Dockerfile."],
            ["AuthMicroservice.Sample",
             "Library-mode demo consumer (SQLite, no email). Exposes /whoami and "
             "/admin-only to prove the wiring end-to-end."],
            ["AuthMicroservice.UnitTests",
             "xUnit + Moq + FluentAssertions: JWT service, refresh rotation, email "
             "templates, request validators."],
            ["AuthMicroservice.IntegrationTests",
             "WebApplicationFactory<Program> + InMemory database: full register / "
             "verify / login / refresh / forgot / reset / change / lockout flow."],
        ],
        col_widths_cm=[6.5, 10],
    )

    # ---------------- 4. Prerequisites ----------------
    add_heading(doc, "4. Prerequisites", level=1)
    add_bullet(doc, ".NET 8 SDK (LTS) installed on the machine.", ())
    add_bullet(doc, "Docker Desktop, if you want to run the standalone microservice via "
                    "docker-compose.", ())
    add_bullet(doc, "A reachable database instance for production use: SQL Server, "
                    "PostgreSQL, or SQLite. For local development the sample uses SQLite "
                    "and the docker-compose stack ships SQL Server 2022.", ())
    add_bullet(doc, "An SMTP server for verification and reset emails. The docker-compose "
                    "override starts Mailhog on port 8025 so you can inspect messages "
                    "from a browser without a real SMTP account.", ())

    # ---------------- 5. Quick start ----------------
    add_heading(doc, "5. Quick Start", level=1)

    add_heading(doc, "5.1 Standalone via Docker Compose", level=2)
    add_para(
        doc,
        "This is the fastest way to see the service running end-to-end. It starts SQL "
        "Server 2022 and Mailhog alongside the API container.",
    )
    add_code_block(doc, """copy .env.example .env
# edit .env - set a real JWT_KEY (>= 32 chars)

docker compose up -d --build

# health
curl http://localhost:8080/auth/health

# swagger UI
Start-Process http://localhost:8080/swagger

# mailhog UI (dev override)
Start-Process http://localhost:8025
""")

    add_heading(doc, "5.2 Library-mode sample", level=2)
    add_para(
        doc,
        "The Sample project references AuthMicroservice.Core and runs on SQLite with "
        "email disabled. It exposes /whoami (any authenticated user) and /admin-only "
        "(users in the Admin role) so you can verify the consumer wiring.",
        code_terms=("AuthMicroservice.Core", "/whoami", "/admin-only"),
    )
    add_code_block(doc, """dotnet run --project src/AuthMicroservice.Sample
# opens http://localhost:5100/swagger
""")

    # ---------------- 6. Integrating into your own project ----------------
    add_heading(doc, "6. Integrating Into Your Own Project", level=1)
    add_para(
        doc,
        "Consuming AuthMicroservice from an existing ASP.NET Core 8 Web API takes three "
        "steps: reference the projects, wire up the pipeline, and add a configuration "
        "section.",
    )

    add_heading(doc, "6.1 Reference the projects", level=2)
    add_para(
        doc,
        "Add project references to AuthMicroservice.Core and to the migration assemblies "
        "for the database providers you plan to support. Most consumers only need one "
        "provider assembly.",
        code_terms=("AuthMicroservice.Core",),
    )

    add_heading(doc, "6.2 Wire up Program.cs", level=2)
    add_code_block(doc, """using AuthMicroservice.Core.Extensions;

var builder = WebApplication.CreateBuilder(args);
builder.Services.AddAuthMicroservice(builder.Configuration);

var app = builder.Build();
app.UseAuthMicroservice();
app.MapAuthMicroservice();
await app.ApplyAuthMicroserviceMigrationsAsync();
app.Run();
""")

    add_para(doc, "The four extension methods are:")
    add_table(
        doc,
        headers=["Method", "Responsibility"],
        rows=[
            ["AddAuthMicroservice(IConfiguration, Action<AuthMicroserviceOptions>?)",
             "Registers EF Core, ASP.NET Identity, JWT bearer authentication, the email "
             "services, validators, and the option types."],
            ["UseAuthMicroservice()",
             "Enables Swagger (when configured), plus authentication and authorization "
             "middleware, in the correct order."],
            ["MapAuthMicroservice()",
             "Maps every /auth/* endpoint under the configured RoutePrefix."],
            ["ApplyAuthMicroserviceMigrationsAsync()",
             "When Database.AutoMigrate is true, runs pending EF Core migrations and "
             "seeds the Admin and User roles."],
        ],
        col_widths_cm=[6.5, 10],
    )

    add_heading(doc, "6.3 Minimum appsettings.json", level=2)
    add_code_block(doc, """{
  "AuthMicroservice": {
    "Database": {
      "Provider": "Sqlite",
      "ConnectionString": "Data Source=auth.db",
      "AutoMigrate": true,
      "SeedDefaults": true
    },
    "Jwt": {
      "Issuer": "MyApp",
      "Audience": "MyApp.Clients",
      "Key": "replace-me-with-a-32-plus-character-secret"
    },
    "Email": { "Enabled": false }
  }
}
""")

    # ---------------- 7. Configuration reference ----------------
    add_heading(doc, "7. Configuration Reference", level=1)
    add_para(
        doc,
        "All settings live under a single AuthMicroservice section. Bind it from any "
        "IConfiguration (typically appsettings.json plus environment variables). The full "
        "shape is shown below, followed by per-field descriptions.",
        code_terms=("AuthMicroservice",),
    )
    add_code_block(doc, """{
  "AuthMicroservice": {
    "Database": {
      "Provider": "SqlServer",       // SqlServer | Postgres | Sqlite | InMemory
      "ConnectionString": "...",
      "AutoMigrate": true,
      "SeedDefaults": true             // seeds Admin + User roles
    },
    "Jwt": {
      "Issuer": "AuthMicroservice",
      "Audience": "AuthMicroservice.Clients",
      "Key": "REPLACE_FROM_ENV_MIN_32_CHARS",
      "AccessTokenLifetimeMinutes": 15,
      "RefreshTokenLifetimeDays": 7,
      "ClockSkewSeconds": 30
    },
    "Email": {
      "Enabled": true,
      "FromAddress": "no-reply@example.com",
      "FromName": "Auth Service",
      "Smtp": {
        "Host": "localhost", "Port": 1025,
        "UseStartTls": false, "UseSsl": false,
        "Username": "", "Password": ""
      },
      "Templates": {
        "VerifyEmailSubject": "Verify your email",
        "PasswordResetSubject": "Reset your password"
      }
    },
    "Identity": {
      "Password": {
        "RequiredLength": 8, "RequireDigit": true,
        "RequireLowercase": true, "RequireUppercase": true,
        "RequireNonAlphanumeric": true, "RequiredUniqueChars": 1
      },
      "Lockout": {
        "AllowedForNewUsers": true,
        "MaxFailedAccessAttempts": 5, "DefaultLockoutMinutes": 15
      },
      "SignIn": {
        "RequireConfirmedEmail": true,
        "RequireConfirmedPhoneNumber": false
      },
      "User": { "RequireUniqueEmail": true }
    },
    "TokenLinks": {
      "EmailVerificationBaseUrl": "https://app.example.com/verify-email",
      "PasswordResetBaseUrl": "https://app.example.com/reset-password"
    },
    "RoutePrefix": "/auth",
    "EnableSwagger": true
  }
}
""")

    add_heading(doc, "7.1 Environment variables", level=2)
    add_para(
        doc,
        "Any nested key can be overridden with an environment variable using the "
        "double-underscore separator. This is the recommended way to inject secrets:",
    )
    add_code_block(doc, """AuthMicroservice__Jwt__Key
AuthMicroservice__Database__ConnectionString
AuthMicroservice__Email__Smtp__Password
""")

    add_heading(doc, "7.2 Fail-fast validation", level=2)
    add_para(doc, "Startup aborts with a clear message when:")
    add_bullet(doc, "Jwt.Key is under 32 characters.", ("Jwt.Key",))
    add_bullet(doc, "Database.Provider is not one of SqlServer, Postgres, Sqlite, or "
                    "InMemory.", ("Database.Provider",))
    add_bullet(doc, "Email.Enabled is true but Email.Smtp.Host is empty.",
               ("Email.Enabled", "Email.Smtp.Host"))

    # ---------------- 8. API reference ----------------
    add_heading(doc, "8. API Reference", level=1)
    add_para(
        doc,
        "All endpoints are mounted under the configured RoutePrefix (default /auth). "
        "Errors are returned as RFC 7807 ProblemDetails with error codes such as "
        "INVALID_CREDENTIALS, USER_LOCKED_OUT, and INVALID_REFRESH_TOKEN.",
        code_terms=("RoutePrefix", "/auth"),
    )

    add_heading(doc, "8.1 Endpoint summary", level=2)
    add_table(
        doc,
        headers=["Method", "Route", "Auth", "Notes"],
        rows=[
            ["POST", "/auth/register", "anon",
             "201 with tokens; sends verification email"],
            ["POST", "/auth/login", "anon",
             "200 tokens; 401 bad creds; 403 unconfirmed; 423 lockout"],
            ["POST", "/auth/refresh", "anon",
             "Rotates refresh; replay of old refresh returns 401"],
            ["POST", "/auth/logout", "auth", "Revokes one refresh token"],
            ["POST", "/auth/logout-all", "auth",
             "Revokes all refresh tokens for the user"],
            ["POST", "/auth/verify-email", "anon", "Body: { userId, token }"],
            ["POST", "/auth/resend-verification", "anon",
             "Silent on unknown email"],
            ["POST", "/auth/forgot-password", "anon",
             "Always 200 (guards against enumeration)"],
            ["POST", "/auth/reset-password", "anon",
             "Body: { email, token, newPassword }"],
            ["POST", "/auth/change-password", "auth",
             "Requires current password"],
            ["GET", "/auth/me", "auth",
             "User profile with roles and claims"],
            ["GET", "/auth/health", "anon", "Liveness probe"],
        ],
        col_widths_cm=[1.8, 4.5, 1.5, 8.2],
    )

    add_heading(doc, "8.2 Request payloads", level=2)
    add_table(
        doc,
        headers=["DTO", "Fields"],
        rows=[
            ["RegisterRequest", "Email (string), Password (string), FullName (string?)"],
            ["LoginRequest", "Email (string), Password (string)"],
            ["RefreshRequest", "AccessToken (string), RefreshToken (string)"],
            ["ChangePasswordRequest", "CurrentPassword (string), NewPassword (string)"],
            ["ResetPasswordRequest",
             "Email (string), Token (string), NewPassword (string)"],
            ["ForgotPasswordRequest", "Email (string)"],
            ["VerifyEmailRequest", "UserId (string), Token (string)"],
        ],
        col_widths_cm=[5.5, 11],
    )

    add_heading(doc, "8.3 Response payloads", level=2)
    add_table(
        doc,
        headers=["DTO", "Fields"],
        rows=[
            ["AuthResponse",
             "AccessToken (string), RefreshToken (string), ExpiresAt (DateTime), "
             "TokenType (string, default 'Bearer'), User (UserResponse)"],
            ["UserResponse",
             "Id (Guid), Email (string), FullName (string?), EmailConfirmed (bool), "
             "Roles (IReadOnlyList<string>), Claims (IReadOnlyDictionary<string, string>)"],
            ["MessageResponse", "Message (string)"],
        ],
        col_widths_cm=[5.5, 11],
    )

    add_heading(doc, "8.4 Sample requests (curl)", level=2)
    add_code_block(doc, """# Register
curl -X POST http://localhost:8080/auth/register \\
  -H "Content-Type: application/json" \\
  -d '{"email":"alice@example.com","password":"P@ssw0rd!","fullName":"Alice"}'

# Login
curl -X POST http://localhost:8080/auth/login \\
  -H "Content-Type: application/json" \\
  -d '{"email":"alice@example.com","password":"P@ssw0rd!"}'

# Refresh
curl -X POST http://localhost:8080/auth/refresh \\
  -H "Content-Type: application/json" \\
  -d '{"accessToken":"<JWT>","refreshToken":"<RT>"}'

# Protected profile
curl http://localhost:8080/auth/me -H "Authorization: Bearer <JWT>"
""")

    # ---------------- 9. End-to-end walkthrough ----------------
    add_heading(doc, "9. End-to-End Walkthrough (PowerShell)", level=1)
    add_para(
        doc,
        "The following sequence exercises the complete flow against a locally running "
        "instance. When Mailhog is used, verification and reset tokens can be picked up "
        "from the Mailhog UI on port 8025.",
    )
    add_code_block(doc, """# 1. Register
$reg = @{ email="alice@example.com"; password="P@ssw0rd!"; fullName="Alice" } |
    ConvertTo-Json
$response = curl -X POST http://localhost:8080/auth/register `
    -H "Content-Type: application/json" -d $reg | ConvertFrom-Json

# 2. Grab the verification link from Mailhog (http://localhost:8025), then:
$ver = @{ userId="<GUID>"; token="<TOKEN>" } | ConvertTo-Json
curl -X POST http://localhost:8080/auth/verify-email `
    -H "Content-Type: application/json" -d $ver

# 3. Login
$login = @{ email="alice@example.com"; password="P@ssw0rd!" } | ConvertTo-Json
$auth = curl -X POST http://localhost:8080/auth/login `
    -H "Content-Type: application/json" -d $login | ConvertFrom-Json

# 4. Protected endpoint
curl http://localhost:8080/auth/me `
    -H "Authorization: Bearer $($auth.accessToken)"

# 5. Rotate refresh token (old refresh becomes invalid)
$rf = @{ accessToken=$auth.accessToken; refreshToken=$auth.refreshToken } |
    ConvertTo-Json
curl -X POST http://localhost:8080/auth/refresh `
    -H "Content-Type: application/json" -d $rf

# 6. Forgot + reset password (token from Mailhog)
curl -X POST http://localhost:8080/auth/forgot-password `
    -H "Content-Type: application/json" `
    -d (@{ email="alice@example.com" } | ConvertTo-Json)

# 7. Trigger lockout (default 5 failed attempts)
1..5 | ForEach-Object {
    curl -X POST http://localhost:8080/auth/login `
        -H "Content-Type: application/json" `
        -d (@{ email="alice@example.com"; password="wrong" } | ConvertTo-Json)
}
""")

    # ---------------- 10. Extending the service ----------------
    add_heading(doc, "10. Extending the Service", level=1)

    add_heading(doc, "10.1 Swap the email transport", level=2)
    add_para(
        doc,
        "The default IEmailSender implementation is SmtpEmailSender (MailKit). To send "
        "through a different provider, register your own implementation after "
        "AddAuthMicroservice. The higher-level IEmailService (which composes verification "
        "and reset templates) is unaffected and will call your implementation directly.",
        code_terms=("IEmailSender", "SmtpEmailSender", "AddAuthMicroservice",
                    "IEmailService"),
    )
    add_code_block(doc, """builder.Services.AddAuthMicroservice(builder.Configuration);
builder.Services.AddSingleton<IEmailSender, MySendGridEmailSender>();
""")

    add_heading(doc, "10.2 Roles and claims", level=2)
    add_para(
        doc,
        "Because ASP.NET Core Identity is used underneath, standard RoleManager and "
        "UserManager APIs are available. Two roles are seeded on first run when "
        "SeedDefaults is true: Admin and User. Add extra roles or claims with the "
        "regular Identity APIs, and enforce them in your own endpoints with the built-in "
        "policies (for example [Authorize(Roles = 'Admin')]).",
        code_terms=("RoleManager", "UserManager", "SeedDefaults",
                    "[Authorize(Roles = 'Admin')]"),
    )

    add_heading(doc, "10.3 Choose a database provider", level=2)
    add_para(
        doc,
        "The provider is chosen through AuthMicroservice.Database.Provider. Reference the "
        "matching migration assembly from your consumer so the migrations can be applied "
        "on startup. Supported values are SqlServer, Postgres, Sqlite, and InMemory.",
        code_terms=("AuthMicroservice.Database.Provider",),
    )

    # ---------------- 11. Database migrations ----------------
    add_heading(doc, "11. Database Migrations", level=1)
    add_para(
        doc,
        "Migrations are split per provider. When a schema change is required, generate a "
        "new migration in every provider assembly that ships with your build. The "
        "migration name should match across providers so the history stays legible.",
    )
    add_code_block(doc, """# SqlServer
dotnet ef migrations add AddSomething `
  --project src/AuthMicroservice.Migrations.SqlServer `
  --startup-project src/AuthMicroservice.Migrations.SqlServer `
  --context AuthDbContext

# Postgres
dotnet ef migrations add AddSomething `
  --project src/AuthMicroservice.Migrations.Postgres `
  --startup-project src/AuthMicroservice.Migrations.Postgres `
  --context AuthDbContext

# Sqlite
dotnet ef migrations add AddSomething `
  --project src/AuthMicroservice.Migrations.Sqlite `
  --startup-project src/AuthMicroservice.Migrations.Sqlite `
  --context AuthDbContext
""")

    # ---------------- 12. Testing ----------------
    add_heading(doc, "12. Testing", level=1)
    add_para(doc, "Run the full test suite:", ())
    add_code_block(doc, "dotnet test\n")
    add_para(doc, "The suite is split into two projects:")
    add_bullet(
        doc,
        "Unit tests: JWT service (claims, expiry, signature), refresh token rotation "
        "(hash storage, chain revoke), email templating, and request validators.",
    )
    add_bullet(
        doc,
        "Integration tests: WebApplicationFactory<Program> with the InMemory provider. "
        "Cover register / login (wrong password returns 401, lockout returns 423), "
        "refresh rotation and replay rejection, forgot / reset flow, and change-password "
        "revoking existing refresh tokens.",
    )

    # ---------------- 13. Deployment notes ----------------
    add_heading(doc, "13. Deployment Notes", level=1)
    add_para(
        doc,
        "The production image is built from src/AuthMicroservice.Api/Dockerfile. It uses "
        "the ASP.NET 8 runtime base image, listens on port 8080, and runs as a non-root "
        "user. docker-compose.yml provisions a SQL Server 2022 container alongside the "
        "API. The docker-compose.override.yml adds Mailhog for local development.",
        code_terms=("src/AuthMicroservice.Api/Dockerfile", "docker-compose.yml",
                    "docker-compose.override.yml"),
    )
    add_para(doc, "Before deploying to any non-development environment, set at least:")
    add_bullet(doc, "AuthMicroservice__Jwt__Key - 32 or more characters, rotated on any "
                    "suspected incident.", ("AuthMicroservice__Jwt__Key",))
    add_bullet(doc, "AuthMicroservice__Database__ConnectionString - point at your real "
                    "database.", ("AuthMicroservice__Database__ConnectionString",))
    add_bullet(doc, "AuthMicroservice__Email__Smtp__Host / Port / Username / Password - "
                    "point at your real SMTP relay when Email.Enabled is true.",
               ("AuthMicroservice__Email__Smtp__Host",))
    add_bullet(doc, "AuthMicroservice__TokenLinks__EmailVerificationBaseUrl and "
                    "PasswordResetBaseUrl - the public URLs of your front-end pages that "
                    "receive the tokens.",
               ("AuthMicroservice__TokenLinks__EmailVerificationBaseUrl",
                "PasswordResetBaseUrl"))

    # ---------------- 14. Troubleshooting ----------------
    add_heading(doc, "14. Troubleshooting", level=1)
    add_table(
        doc,
        headers=["Symptom", "Likely cause and fix"],
        rows=[
            ["Startup exits with 'Jwt.Key must be at least 32 characters'.",
             "Set AuthMicroservice__Jwt__Key (env var) or Jwt.Key (config) to a longer "
             "secret. Rotate immediately if the previous key leaked."],
            ["Startup exits with 'Unknown database provider'.",
             "Database.Provider must be one of SqlServer, Postgres, Sqlite, InMemory. "
             "Check case-sensitivity and the assembly reference for the chosen provider."],
            ["Startup exits with 'Email.Smtp.Host is required'.",
             "Either set Email.Enabled to false for local dev, or fill in Email.Smtp.Host "
             "(and credentials if the relay requires them)."],
            ["/auth/login returns 403.",
             "The account has not verified its email. Trigger /auth/resend-verification "
             "or set Identity.SignIn.RequireConfirmedEmail to false in development."],
            ["/auth/login returns 423.",
             "The account is locked after too many failed attempts. Wait for "
             "Identity.Lockout.DefaultLockoutMinutes, or unlock the user through your "
             "own admin tooling."],
            ["/auth/refresh returns 401 on a token that used to work.",
             "The refresh token was already rotated (its replacement should be used) or "
             "revoked because a replay was detected. Log in again."],
        ],
        col_widths_cm=[5.5, 11],
    )

    # ---------------- 15. Roadmap ----------------
    add_heading(doc, "15. Roadmap", level=1)
    add_para(doc, "Under consideration for future releases:")
    add_bullet(doc, "Two-factor authentication (TOTP).", ())
    add_bullet(doc, "External OAuth providers - Google and Microsoft first.", ())
    add_bullet(doc, "Rate limiting on /auth/login and /auth/forgot-password.", ())
    add_bullet(doc, "Audit log of authentication events.", ())
    add_bullet(doc, "Health checks for the database and SMTP relay.", ())
    add_bullet(doc, "OpenTelemetry traces and structured logging with correlation IDs.", ())
    add_bullet(doc, "Publishing AuthMicroservice.Core as a NuGet package so downstream "
                    "projects can consume it as a dependency instead of a project "
                    "reference.", ("AuthMicroservice.Core",))

    # ---------------- 16. License ----------------
    add_heading(doc, "16. License", level=1)
    add_para(
        doc,
        "AuthMicroservice is released under the MIT License. It can be freely used as "
        "the authentication foundation of internal, commercial, or open-source projects.",
    )

    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT} ({OUTPUT.stat().st_size:,} bytes)")


if __name__ == "__main__":
    build_document()
